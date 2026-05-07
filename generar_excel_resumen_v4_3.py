import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils.dataframe import dataframe_to_rows

# Cargar datos
v43 = pd.read_csv('backend/data/processed/imputation/Datos_Imputados_v4_3_Corrected.csv')
models = pd.read_csv('backend/data/processed/imputation/Tabla_Modelos_HH_v4_3.csv')

# ============================================================
# HOJA 1: RESUMEN POR ESTACION
# ============================================================
resumen = v43.groupby(['NombreEstacion', 'Nivel_Imputacion']).size().unstack(fill_value=0)
resumen['Total'] = resumen.sum(axis=1)
for col in [0, 1, 2, 3]:
    if col not in resumen.columns:
        resumen[col] = 0
    resumen[f'P_{col}'] = (resumen[col] / resumen['Total'] * 100).round(2)

resumen = resumen[[0, 1, 2, 3, 'Total', 'P_0', 'P_1', 'P_2', 'P_3']].reset_index()
resumen.columns = ['Estacion', 'L0_Original', 'L1_Interpolacion', 'L2_HH', 'L3_missForest', 
                   'Total_Registros', 'Pct_Original', 'Pct_L1', 'Pct_L2', 'Pct_L3']

# Ordenar por orden geografico (aguas abajo a arriba)
geo_order = ['San Pedrito', 'Calamar', 'Tenerife', 'Plato', 'Magangue', 'Barbosa', 'Coyongal',
             'Armenia', 'Rio Nuevo', 'Las Palomas', 'Las Aguadas', 'El Banco', 'Penoncito',
             'La Gloria', 'Gamarra Aut', 'El Contento', 'Sitio Nuevo', 'Puerto Wilches']
resumen['orden'] = resumen['Estacion'].map({v: i for i, v in enumerate(geo_order)})
resumen = resumen.sort_values('orden').drop(columns=['orden'])

# ============================================================
# HOJA 2: MODELOS H-H
# ============================================================
models_sheet = models[['Recipiente', 'Independiente', 'Ecuacion', 'R2', 'r', 'N_Entrenamiento']].copy()
models_sheet.columns = ['Estacion_Dependiente', 'Estacion_Independiente', 'Ecuacion_Cuadratica', 'R2', 'r', 'N_Entrenamiento']

# ============================================================
# CREAR EXCEL
# ============================================================
wb = Workbook()

# --- HOJA 1: Resumen ---
ws1 = wb.active
ws1.title = "Resumen por Estacion"

# Titulo
ws1.merge_cells('A1:J1')
ws1['A1'] = "RESUMEN DE IMPUTACION v4.3 - Rio Magdalena"
ws1['A1'].font = Font(bold=True, size=14)
ws1['A1'].alignment = Alignment(horizontal='center')

# Subtitulo
ws1.merge_cells('A2:J2')
ws1['A2'] = "Corte: 31-12-2025 | Hard rule aplicada | Orientacion H-H corregida (aguas arriba -> aguas abajo)"
ws1['A2'].alignment = Alignment(horizontal='center')

# Encabezados
headers = list(resumen.columns)
for col_idx, header in enumerate(headers, 1):
    cell = ws1.cell(row=4, column=col_idx, value=header)
    cell.font = Font(bold=True, color="FFFFFF")
    cell.fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
    cell.alignment = Alignment(horizontal='center')

# Datos
for r_idx, row in enumerate(resumen.itertuples(index=False), 5):
    for c_idx, value in enumerate(row, 1):
        cell = ws1.cell(row=r_idx, column=c_idx, value=value)
        cell.alignment = Alignment(horizontal='center')
        if c_idx == 1:  # Nombre estacion
            cell.alignment = Alignment(horizontal='left')

# Ajustar anchos
ws1.column_dimensions['A'].width = 18
for col in ['B', 'C', 'D', 'E', 'F']:
    ws1.column_dimensions[col].width = 14
for col in ['G', 'H', 'I', 'J']:
    ws1.column_dimensions[col].width = 12

# --- HOJA 2: Modelos H-H ---
ws2 = wb.create_sheet(title="Modelos H-H v4.3")

ws2.merge_cells('A1:F1')
ws2['A1'] = "MODELOS ALTURA-ALTURA (H-H) v4.3 - Orientacion hidrologica correcta"
ws2['A1'].font = Font(bold=True, size=14)
ws2['A1'].alignment = Alignment(horizontal='center')

ws2.merge_cells('A2:F2')
ws2['A2'] = "Criterio: Independiente (aguas arriba) -> Dependiente (aguas abajo) | R2 >= 0.75"
ws2['A2'].alignment = Alignment(horizontal='center')

headers2 = list(models_sheet.columns)
for col_idx, header in enumerate(headers2, 1):
    cell = ws2.cell(row=4, column=col_idx, value=header)
    cell.font = Font(bold=True, color="FFFFFF")
    cell.fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
    cell.alignment = Alignment(horizontal='center')

for r_idx, row in enumerate(models_sheet.itertuples(index=False), 5):
    for c_idx, value in enumerate(row, 1):
        cell = ws2.cell(row=r_idx, column=c_idx, value=value)
        cell.alignment = Alignment(horizontal='center')
        if c_idx == 3:  # Ecuacion
            cell.alignment = Alignment(horizontal='left')

ws2.column_dimensions['A'].width = 20
ws2.column_dimensions['B'].width = 24
ws2.column_dimensions['C'].width = 45
ws2.column_dimensions['D'].width = 10
ws2.column_dimensions['E'].width = 10
ws2.column_dimensions['F'].width = 16

# Guardar
output_path = 'backend/data/processed/imputation/Resumen_Imputacion_v4_3.xlsx'
wb.save(output_path)
print(f"Excel guardado: {output_path}")
print(f"\nHojas creadas:")
print(f"  1. Resumen por Estacion: {len(resumen)} estaciones")
print(f"  2. Modelos H-H v4.3: {len(models_sheet)} modelos")

# Mostrar resumen en consola
print("\n" + "="*80)
print("RESUMEN POR ESTACION")
print("="*80)
print(resumen.to_string(index=False))

print("\n" + "="*80)
print("TOTALES GLOBALES")
print("="*80)
total_registros = resumen['Total_Registros'].sum()
total_l0 = resumen['L0_Original'].sum()
total_l1 = resumen['L1_Interpolacion'].sum()
total_l2 = resumen['L2_HH'].sum()
total_l3 = resumen['L3_missForest'].sum()
print(f"Total registros: {total_registros:,}")
print(f"L0 Original:     {total_l0:,} ({total_l0/total_registros*100:.2f}%)")
print(f"L1 Interpolacion:{total_l1:,} ({total_l1/total_registros*100:.2f}%)")
print(f"L2 H-H:          {total_l2:,} ({total_l2/total_registros*100:.2f}%)")
print(f"L3 missForest:   {total_l3:,} ({total_l3/total_registros*100:.2f}%)")
