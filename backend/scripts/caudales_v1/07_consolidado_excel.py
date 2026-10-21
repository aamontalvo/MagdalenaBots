#!/usr/bin/env python3
"""
07_consolidado_excel.py
=======================

Genera el consolidado de resumen para revision tecnica:

- backend/entregables/caudales_v1/Consolidado_Caudales_Niveles.xlsx
- backend/entregables/caudales_v1/README.md

El Excel contiene resumenes, no las series completas. Las series viven en
los CSV de origen y se referencian en la hoja Fuentes.
"""

import sys
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

sys.stdout.reconfigure(encoding="utf-8", errors="replace")

ROOT = Path(__file__).resolve().parents[3]
OUT_DIR = ROOT / "backend" / "entregables" / "caudales_v1"
OUT_DIR.mkdir(parents=True, exist_ok=True)
OUT_XLSX = OUT_DIR / "Consolidado_Caudales_Niveles.xlsx"
OUT_README = OUT_DIR / "README.md"

NIV_SUM = ROOT / "backend" / "entregables" / "niveles_v2_1" / "05_validacion_y_auditoria" / "Resumen_Imputacion.csv"
CAUD_SUM = ROOT / "backend" / "data" / "work" / "caudales_v1" / "output" / "Resumen_Imputacion_Caudales.csv"
NC = ROOT / "cartillas_estaciones" / "Caudales" / "curvas_nivel_caudal_resumen.csv"
DUR = ROOT / "cartillas_estaciones" / "Caudales" / "curvas_duracion_caudales_resumen.csv"
HOMOG = ROOT / "backend" / "data" / "processed" / "Resultados_Pruebas_Homogeneidad.csv"

niv = pd.read_csv(NIV_SUM)
caud = pd.read_csv(CAUD_SUM)
nc = pd.read_csv(NC)
dur = pd.read_csv(DUR)
homog = pd.read_csv(HOMOG, sep=";", encoding="utf-8-sig")

kpis = pd.DataFrame(
    [
        ("Niveles - registros finales", 394772, "Datos_Imputados.csv"),
        ("Niveles - L0 validos", 335501, "84,99%"),
        ("Niveles - L1", 1222, "0,31%"),
        ("Niveles - L2", 36758, "9,31%"),
        ("Niveles - L3", 16855, "4,27%"),
        ("Niveles - total imputado", 54835, "13,89%"),
        ("Niveles - NaN", 4436, "1,12%"),
        ("Niveles - cobertura media", 98.87, "%"),
        ("Niveles - estaciones", 24, "24/24 en el informe"),
        ("Niveles - figuras", 150, "143 cartillas + 7 globales"),
        ("Caudales - registros finales", 299239, "Datos_Imputados_Caudales.csv"),
        ("Caudales - L0 validos", 262425, "87,70%"),
        ("Caudales - L1", 349, "0,12%"),
        ("Caudales - L2", 28632, "9,57%"),
        ("Caudales - total imputado", 28981, "9,68%"),
        ("Caudales - NaN", 7833, "2,62%"),
        ("Caudales - cobertura media", 97.74, "%"),
        ("Caudales - estaciones finales", 14, "14 comunes con niveles"),
        ("Caudales - figuras", 101, "98 cartillas + 3 globales"),
        ("Caudales - pares N-C", 222993, "14 modelos"),
        ("Caudales - R2 N-C minimo", 0.6734, "Barbosa"),
        ("Caudales - R2 N-C maximo", 0.9676, "Las Aguadas"),
        ("Universo unico", 35, "14 comunes + 10 solo-niveles + 11 solo-caudales"),
        ("Cartillas totales", 38, "24 niveles + 14 caudales"),
        ("Figuras vigentes totales", 251, "150 niveles + 101 caudales"),
    ],
    columns=["Metrica", "Valor", "Nota"],
)

estaciones = pd.DataFrame(
    [
        ("Niveles", "Incluidas", 24, "24/24; 4 marcadas como inhomogeneas"),
        ("Caudales", "Incluidas", 11, "11/14 comunes"),
        ("Caudales", "Faltantes", 3, "San Pablo, Puerto Berrio, Puerto Salgar"),
        ("Caudales", "No incluidas", 14, "3 comunes faltantes + 11 solo-caudales"),
        ("Universo", "Unicas", 35, "14 comunes + 10 solo-niveles + 11 solo-caudales"),
        ("Universo", "Solo-caudales", 11, "Sin niveles"),
        ("Universo", "Solo-niveles", 10, "Sin caudales"),
    ],
    columns=["Conjunto", "Categoria", "Cantidad", "Detalle"],
)

figuras = pd.DataFrame(
    [
        ("Niveles", "Cartillas", 24, "Una por estacion"),
        ("Niveles", "PNG de cartilla", 143, "Fig1-Fig6; Puerto Salgar sin Fig6"),
        ("Niveles", "PNG globales", 7, "Boxplot, correlacion, heatmap, etc."),
        ("Niveles", "Total", 150, "Vigente"),
        ("Caudales", "Cartillas", 14, "Una por estacion"),
        ("Caudales", "PNG de cartilla", 98, "Fig1-Fig6; Fig5 max/min"),
        ("Caudales", "PNG globales", 3, "Boxplot, correlacion, heatmap"),
        ("Caudales", "Total", 101, "Vigente"),
        ("Total", "Figuras vigentes", 251, "150 + 101"),
    ],
    columns=["Conjunto", "Tipo", "Cantidad", "Detalle"],
)

dudas = pd.DataFrame(
    [
        ("Estadistico Cramer", "Descartado", "No se reproduce", "Documentar que no se uso"),
        ("El Banco sub-diario vs equipo", "Abierto", "Pettitt/MK no coinciden", "Confirmar version de datos"),
        ("Barbosa sin L2", "Abierto", "Cobertura 83,22%", "Evaluar L3 Q-Q o N-C segmentada"),
        ("Puerto Salgar 1937-1960", "Abierto", "Sin niveles concurrentes", "Mantener NaN documentado"),
        ("L3 Q-Q entre estaciones", "Standby", "Barbosa y otras <95%", "Disenar y ejecutar"),
        ("Curvas N-C segmentadas", "Standby", "Barbosa y Armenia", "Evaluar por periodo"),
        ("11 vs 14 estaciones en el informe", "Abierto", "Faltan 3 comunes", "Actualizar el borrador"),
        ("Outliers vs equipo", "Resuelto", "Diferencias de pruebas", "Base oficial sin outliers"),
        ("29-feb y agregacion sub-diaria", "Resuelto", "Evita sesgos", "Documentado en scripts"),
        ("Figuras sin caption en caudales", "Abierto", "20 imagenes", "Agregar captions"),
        ("Fig6 de Puerto Salgar en niveles", "Abierto", "Sin modelo H-H", "Nota de no aplicabilidad"),
        ("Versionado de datos y figuras", "Abierto", "Fuera de git", "Release/ZIP + manifiesto"),
    ],
    columns=["Tema", "Estado", "Impacto", "Siguiente paso"],
)

reportes = pd.DataFrame(
    [
        ("Caudales - mayo 2026", 11, 37, 38, "Historico"),
        ("Caudales - julio 2026", 11, 57, 37, "Borrador vigente; 20 imagenes sin caption"),
        ("Niveles - julio 2026", 24, 105, 112, "Borrador vigente; faltan Fig5/Fig6"),
    ],
    columns=["Borrador", "Estaciones", "Imagenes", "Leyendas", "Estado"],
)

fuentes = pd.DataFrame(
    [
        ("Niveles crudo", "backend/entregables/niveles_v2_1/01_datos_originales/Resumen_Crudos_IDEAM_2026_v2_1.csv"),
        ("Niveles final", "backend/entregables/niveles_v2_1/02_datos_imputados/Datos_Imputados.csv"),
        ("Niveles resumen", "backend/entregables/niveles_v2_1/05_validacion_y_auditoria/Resumen_Imputacion.csv"),
        ("Caudales crudo largo", "backend/data/raw/SERIE_CAUDALES/analisis/caudales_formato_largo.csv"),
        ("Caudales unificado", "backend/data/processed/caudales_unificados.csv"),
        ("Caudales final", "backend/data/work/caudales_v1/output/Datos_Imputados_Caudales.csv"),
        ("Caudales resumen", "backend/data/work/caudales_v1/output/Resumen_Imputacion_Caudales.csv"),
        ("Curvas N-C", "cartillas_estaciones/Caudales/curvas_nivel_caudal_resumen.csv"),
        ("Curvas duracion", "cartillas_estaciones/Caudales/curvas_duracion_caudales_resumen.csv"),
        ("Homogeneidad", "backend/data/processed/Resultados_Pruebas_Homogeneidad.csv"),
        ("Verificacion caudales", "backend/data/raw/SERIE_CAUDALES/analisis/verificacion_caudales.txt"),
        ("Metricas", "docs/METRICAS_CAUDALES_NIVELES.md"),
        ("Contexto", "docs/CONTEXTO_MAESTRO_CAUDALES.md"),
        ("Comentarios Stefany", "docs/RESPUESTAS_COMENTARIOS_STEFANY.md"),
    ],
    columns=["Recurso", "Ruta"],
)

with pd.ExcelWriter(OUT_XLSX, engine="openpyxl") as writer:
    kpis.to_excel(writer, sheet_name="KPIs", index=False)
    niv.to_excel(writer, sheet_name="Niveles_Resumen", index=False)
    caud.to_excel(writer, sheet_name="Caudales_Resumen", index=False)
    nc.to_excel(writer, sheet_name="Curvas_NC", index=False)
    dur.to_excel(writer, sheet_name="Curvas_Duracion", index=False)
    homog.to_excel(writer, sheet_name="Homogeneidad", index=False)
    estaciones.to_excel(writer, sheet_name="Estaciones", index=False)
    figuras.to_excel(writer, sheet_name="Figuras", index=False)
    dudas.to_excel(writer, sheet_name="Dudas", index=False)
    reportes.to_excel(writer, sheet_name="Reportes", index=False)
    fuentes.to_excel(writer, sheet_name="Fuentes", index=False)

# Formato
wb = load_workbook(OUT_XLSX)
header_fill = PatternFill("solid", fgColor="1F4E79")
header_font = Font(color="FFFFFF", bold=True)
thin = Side(style="thin", color="D9D9D9")
border = Border(left=thin, right=thin, top=thin, bottom=thin)

for ws in wb.worksheets:
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for row in ws.iter_rows(min_row=1, max_row=ws.max_row, max_col=ws.max_column):
        for cell in row:
            cell.border = border
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    for col in range(1, ws.max_column + 1):
        max_len = 0
        for cell in ws[get_column_letter(col)]:
            if cell.value is not None:
                max_len = max(max_len, len(str(cell.value)))
        ws.column_dimensions[get_column_letter(col)].width = min(max(12, max_len + 2), 60)

# Portada
ws = wb.create_sheet("Portada", 0)
ws["A1"] = "Consolidado de niveles y caudales - Rio Magdalena"
ws["A1"].font = Font(size=16, bold=True, color="1F4E79")
ws["A2"] = "BPIN 2024000100122 | Contrato UN-OJ-2025-66227"
ws["A3"] = "Corte de datos: 24-ago-2026 | Generado: 19-oct-2026"
ws["A5"] = "Contenido"
ws["A5"].font = Font(bold=True)
contenido = [
    "KPIs: tablero ejecutivo de cifras.",
    "Niveles_Resumen: cobertura e imputacion por estacion.",
    "Caudales_Resumen: cobertura, L1/L2 y R2 N-C por estacion.",
    "Curvas_NC: ecuaciones y bondad de ajuste.",
    "Curvas_Duracion: P50/P90/P95/P100 de caudales.",
    "Homogeneidad: pruebas del equipo en 25 estaciones.",
    "Estaciones: incluidas y no incluidas en los informes.",
    "Figuras: conteo de cartillas y PNG.",
    "Dudas: decisiones abiertas de procesamiento.",
    "Reportes: estado de los borradores.",
    "Fuentes: rutas de los archivos originales.",
]
for i, text in enumerate(contenido, start=6):
    ws[f"A{i}"] = text
ws.column_dimensions["A"].width = 95
ws.sheet_view.showGridLines = False

wb.save(OUT_XLSX)

OUT_README.write_text(
    """# Consolidado de niveles y caudales

## Que es

Resumen ejecutivo y tecnico del estado del analisis de niveles y caudales
del Rio Magdalena al corte del 24-ago-2026.

## Archivo principal

- `Consolidado_Caudales_Niveles.xlsx`

Hojas:

1. `Portada`
2. `KPIs`
3. `Niveles_Resumen`
4. `Caudales_Resumen`
5. `Curvas_NC`
6. `Curvas_Duracion`
7. `Homogeneidad`
8. `Estaciones`
9. `Figuras`
10. `Dudas`
11. `Reportes`
12. `Fuentes`

## Que NO contiene

- Las series completas (394.772 filas de niveles y 299.239 de caudales).
- Las 251 figuras.
- Los borradores `.docx`.

Las rutas a esos archivos estan en la hoja `Fuentes`. Para compartir el
paquete completo, agregar las carpetas de datos y figuras indicadas.

## Cifras clave

- Niveles: 394.772 registros; 335.501 L0 (85,0%); 54.835 imputados (13,9%);
  4.436 NaN (1,1%); cobertura media 98,87%.
- Caudales: 299.239 registros; 262.425 L0 (87,7%); 28.981 imputados (9,7%);
  7.833 NaN (2,6%); cobertura media 97,74%.
- 35 estaciones unicas; 38 cartillas; 251 figuras vigentes.

## Documentos de apoyo

- `docs/METRICAS_CAUDALES_NIVELES.md`
- `docs/CONTEXTO_MAESTRO_CAUDALES.md`
- `docs/INVENTARIO_ENTREGABLES.md`
- `docs/ESTADO_CAUDALES_PARA_STEPHANIE.md`
- `docs/RESPUESTAS_COMENTARIOS_STEFANY.md`
""",
    encoding="utf-8",
)

print(f"OK: {OUT_XLSX}")
print(f"OK: {OUT_README}")
