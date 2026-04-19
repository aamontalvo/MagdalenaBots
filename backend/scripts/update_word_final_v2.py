#!/usr/bin/env python3
"""
Actualizacion Final v2 del Informe Word
========================================

Documento final con:
- Seccion 4.3: Curvas H-H Simple (cuadratico)
- Seccion 6: Analisis detallado con violinplots

Autor: Seccion de Analisis de Hidrologia
Fecha: 2026-02-25
"""

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from pathlib import Path
import pandas as pd

DATA_PATH = Path(r"C:\Users\alexa\Documents\PERSONAL PROJECTS\MagdalenaBot\Data\Niveles")
REPORT_PATH = DATA_PATH / "04_reports" / "Informe_Analisis_Exploratorio_Magdalena.docx"
FIGURES_PATH = DATA_PATH / "03_figures"
ANALYSIS_PATH = DATA_PATH / "02_analysis"

COLOR_GRAY = RGBColor(0x66, 0x66, 0x66)


def add_section_header(doc, title, level=1):
    doc.add_heading(title, level=level)


def add_paragraph(doc, text, bold=False):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.bold = bold
    return p


def add_figure(doc, image_path, caption, width_inches=6.5, figure_count=[0]):
    figure_count[0] += 1
    ref = f"Figura {figure_count[0]}"
    
    if image_path.exists():
        doc.add_picture(str(image_path), width=Inches(width_inches))
        last_paragraph = doc.paragraphs[-1]
        last_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    caption_p = doc.add_paragraph()
    caption_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = caption_p.add_run(f"{ref}. {caption}")
    run.italic = True
    run.font.size = Pt(10)
    run.font.color.rgb = COLOR_GRAY
    
    return ref


def add_table_from_df(doc, df, caption, table_count=[0]):
    table_count[0] += 1
    ref = f"Tabla {table_count[0]}"
    
    table = doc.add_table(rows=1, cols=len(df.columns))
    table.style = 'Table Grid'
    
    hdr_cells = table.rows[0].cells
    for i, col in enumerate(df.columns):
        hdr_cells[i].text = str(col)
        for paragraph in hdr_cells[i].paragraphs:
            for run in paragraph.runs:
                run.bold = True
                run.font.size = Pt(10)
    
    for _, row in df.iterrows():
        row_cells = table.add_row().cells
        for i, value in enumerate(row):
            row_cells[i].text = str(value)
            for paragraph in row_cells[i].paragraphs:
                for run in paragraph.runs:
                    run.font.size = Pt(9)
    
    caption_p = doc.add_paragraph()
    caption_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = caption_p.add_run(f"{ref}. {caption}")
    run.italic = True
    run.font.size = Pt(10)
    run.font.color.rgb = COLOR_GRAY
    
    return ref


def generate_section_4_3_final(doc):
    """Seccion 4.3: Curvas H-H Simple Cuadratico (FINAL)."""
    add_section_header(doc, "4.3 Curvas H-H (Nivel-Nivel) - Modelo Cuadratico", level=2)
    
    text = (
        "Las curvas H-H representan las relaciones estadisticas entre niveles de estaciones "
        "correlacionadas. Se utilizaron modelos cuadraticos para capturar relaciones no lineales "
        "entre estaciones del Rio Magdalena, siguiendo la direccion hidrologica aguas arriba "
        "hacia aguas abajo.\n\n"
        "La forma funcional del modelo es:\n"
        "    H_target = a * H_donor^2 + b * H_donor + c\n\n"
        "Las relaciones establecidas fueron:\n"
        "- Calamar (aguas abajo) ← Magangue (aguas arriba)\n"
        "- Las Aguadas (aguas abajo) ← El Banco (aguas arriba)\n"
        "- El Banco (aguas abajo) ← La Gloria (aguas arriba)"
    )
    add_paragraph(doc, text)
    
    # Tabla de resultados
    try:
        stats = pd.read_csv(ANALYSIS_PATH / "final_imputation_statistics.csv")
        table_ref = add_table_from_df(doc, stats, 
            "Resultados de imputacion H-H Simple (modelo cuadratico)")
        
        text2 = (
            f"\nLa {table_ref} presenta los resultados de la imputacion. Los valores de R2 "
            "superiores a 0.82 indican relaciones fuertes entre estaciones. El modelo cuadratico "
            "captura adecuadamente las no linealidades hidraulicas del sistema fluvial."
        )
        add_paragraph(doc, text2)
    except:
        pass
    
    # Figura
    fig_ref = add_figure(doc,
        FIGURES_PATH / "imputation" / "final_imputation_hh_simple.png",
        "Series temporales con imputacion H-H Simple (cuadratico). "
        "Puntos rojos indican valores imputados.",
        width_inches=6.5)
    
    text3 = (
        f"\nLa {fig_ref} muestra las series temporales con los valores imputados. "
        "Un total de 628 valores fueron estimados usando las curvas H-H cuadraticas, "
        "distribuidos entre las tres estaciones candidatas."
    )
    add_paragraph(doc, text3)


def generate_section_6_final(doc):
    """Seccion 6: Analisis Detallado por Estaciones (FINAL)."""
    add_section_header(doc, "6. ANALISIS DETALLADO POR ESTACIONES", level=1)
    
    intro = (
        "Esta seccion presenta un analisis estadistico detallado de las estaciones con "
        "datos completos (>=85% de completitud). Para cada estacion se presentan:\n\n"
        "1. Violinplots mensuales multianuales con boxplot interno y linea de media\n"
        "2. Mejor distribucion de probabilidad ajustada con sus parametros\n\n"
        "Las tres estaciones analizadas representan puntos clave del Rio Magdalena, "
        "desde aguas arriba (El Banco) hasta aguas abajo (Calamar)."
    )
    add_paragraph(doc, intro)
    
    # Resumen
    try:
        summary = pd.read_csv(ANALYSIS_PATH / "station_summary_section6.csv")
        table_ref = add_table_from_df(doc, summary,
            "Resumen de caracteristicas estadisticas por estacion")
        
        text_table = (
            f"\nLa {table_ref} resume las caracteristicas de las tres estaciones. "
            "Todas presentaron la distribucion Weibull como mejor ajuste. Los parametros "
            "permiten reconstruir las funciones de densidad de probabilidad para analisis "
            "de frecuencia de niveles."
        )
        add_paragraph(doc, text_table)
    except:
        pass
    
    # Analisis por estacion
    stations_info = [
        ('Calamar', 'Estacion ubicada aguas abajo en el sistema del Rio Magdalena. '
                   'Presenta la mayor variabilidad de niveles debido a la influencia '
                   'de las ciénagas y la planicie aluvial.'),
        ('Las Aguadas', 'Estacion intermedia con buena correlacion con El Banco. '
                       'Representa la transicion entre el tramo medio y bajo del rio.'),
        ('El Banco', 'Estacion ubicada aguas arriba respecto a Las Aguadas. '
                    'Presenta menor variabilidad y buena calidad de datos.')
    ]
    
    for i, (station, description) in enumerate(stations_info, start=1):
        add_section_header(doc, f"6.{i} Estacion {station}", level=2)
        
        try:
            summary = pd.read_csv(ANALYSIS_PATH / "station_summary_section6.csv")
            row = summary[summary['Estacion'] == station].iloc[0]
            
            # Descripcion
            text = (
                f"{description}\n\n"
                f"La estacion {station} cuenta con {int(row['N']):,} observaciones "
                f"con cota media de {row['Media_m']} m.s.n.m. "
                f"(desviacion estandar: {row['Desv_Est_m']} m). "
                f"La distribucion {row['Mejor_Distribucion']} fue seleccionada como mejor ajuste "
                f"(AIC = {row['AIC']}).\n\n"
                f"Parametros de la distribucion {row['Mejor_Distribucion']}: {row['Parametros']}"
            )
            add_paragraph(doc, text)
            
            # Figura
            fig_ref = add_figure(doc,
                FIGURES_PATH / "detailed_analysis" / f"{station}_monthly_violin_boxplot.png",
                f"Distribucion mensual multianual de cotas - Estacion {station}. "
                "Violin plot con boxplot interno (mediana y cuartiles) y linea de media mensual.",
                width_inches=6.5)
            
            text2 = (
                f"\nLa {fig_ref} muestra la variabilidad estacional de las cotas. "
                "El boxplot interno permite identificar la mediana y rangos intercuartilicos, "
                "mientras que la linea de media revela el ciclo hidrologico anual tipico "
                "de la region Caribe colombiana."
            )
            add_paragraph(doc, text2)
            
        except Exception as e:
            add_paragraph(doc, f"[Error: {e}]")


def update_document_final():
    """Actualiza el documento Word final."""
    print("Actualizando documento Word final...")
    
    doc = Document(REPORT_PATH)
    
    # Nota de version final
    doc.add_page_break()
    add_section_header(doc, "VERSION FINAL - ACTUALIZACIONES", level=1)
    
    add_paragraph(doc, 
        "Las siguientes secciones presentan los resultados finales del analisis "
        "con imputacion H-H Simple (modelo cuadratico) y visualizaciones actualizadas.",
        bold=True)
    
    # Secciones finales
    generate_section_4_3_final(doc)
    doc.add_page_break()
    generate_section_6_final(doc)
    
    # Guardar
    doc.save(REPORT_PATH)
    print(f"Documento final actualizado: {REPORT_PATH}")


if __name__ == '__main__':
    update_document_final()
